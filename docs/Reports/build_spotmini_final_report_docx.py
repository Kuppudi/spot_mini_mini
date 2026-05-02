#!/usr/bin/env python3
"""Build the SpotMini final report DOCX from the markdown source."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "spotmini_final_report.md"
OUTPUT = ROOT / "SpotMini_Final_Report.docx"
IMAGE_RE = re.compile(r"!\[(?P<alt>.*?)\]\((?P<path>.*?)\)")


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.bold = True

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_page_number(footer_p)


def add_paragraph(document: Document, text: str, *, style: str | None = None, center: bool = False, italic: bool = False) -> None:
    paragraph = document.add_paragraph(style=style)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = italic
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15


def add_image(document: Document, alt_text: str, image_path: str) -> None:
    resolved = Path(image_path)
    if not resolved.is_absolute():
        resolved = ROOT / resolved

    if not resolved.exists():
        add_paragraph(
            document,
            f"[Missing image: {image_path}]",
            italic=True,
            center=True,
        )
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(resolved), width=Inches(6.8))
    paragraph.paragraph_format.space_after = Pt(6)

    if alt_text:
        add_paragraph(document, alt_text, center=True, italic=True)


def build_docx() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()

    document = Document()
    configure_document(document)

    on_cover = True
    for raw_line in lines:
        line = raw_line.rstrip()

        if not line:
            document.add_paragraph("")
            continue

        if line == "<PAGEBREAK>":
            document.add_page_break()
            on_cover = False
            continue

        if line.startswith("# "):
            heading = line[2:].strip()
            if on_cover:
                add_paragraph(document, heading, center=True)
            else:
                document.add_heading(heading, level=1)
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            if on_cover:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(heading)
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(13)
            else:
                document.add_heading(heading, level=2)
            continue

        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
            continue

        if line.startswith("- "):
            add_paragraph(document, line[2:].strip(), style="List Bullet")
            continue

        image_match = IMAGE_RE.fullmatch(line.strip())
        if image_match:
            add_image(
                document,
                image_match.group("alt").strip(),
                image_match.group("path").strip(),
            )
            continue

        if line.startswith("[") and line.endswith("]"):
            add_paragraph(document, line, italic=True)
            continue

        add_paragraph(document, line)

    document.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
